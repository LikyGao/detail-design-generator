from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient

from local_backend import desktop
from local_backend.app import APP_HEADER, APP_HEADER_VALUE, create_app

ROOT = Path(__file__).resolve().parents[1]


def _valid_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("株式会社〇〇 様", style="Title")
    document.add_paragraph("サンプルプロジェクト", style="Title")
    document.add_table(rows=4, cols=3)
    document.add_table(rows=2, cols=4)
    document.add_paragraph("1 概要", style="Heading 1")
    document.add_paragraph("概要の標準本文です。", style="Normal")
    document.add_paragraph("1.1 目的", style="Heading 2")
    document.add_paragraph("目的の標準本文です。", style="Normal")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _register_template(client: TestClient, content: bytes | None = None):
    return client.post(
        "/api/templates/register",
        data={"document_type": "server_storage", "template_version": "1.2"},
        files={
            "template_file": (
                "基本設計書_server_v1.2.docx",
                content or _valid_template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )


def test_health_and_application_identity():
    response = TestClient(create_app()).get("/api/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}
    assert response.headers[APP_HEADER] == APP_HEADER_VALUE


def test_root_injects_local_bridge_without_mutating_source():
    response = TestClient(create_app()).get("/")
    assert response.status_code == 200
    source = (ROOT / "基本設計書generator.html").read_text(encoding="utf-8")
    bridge_tag = '<script src="/local-bridge.js"></script>'
    assert bridge_tag not in source
    assert bridge_tag in response.text
    assert response.text.replace("  " + bridge_tag + "\n", "", 1) == source
    assert response.headers[APP_HEADER] == APP_HEADER_VALUE


def test_local_bridge_routes_personal_dify_features_to_local_apis():
    response = TestClient(create_app()).get("/local-bridge.js")
    assert response.status_code == 200
    assert response.headers[APP_HEADER] == APP_HEADER_VALUE
    assert "/api/template-data" in response.text
    assert "/api/templates/register" in response.text
    assert "/api/generate-word" in response.text
    assert "difyCall = async function localAwareDifyCall" in response.text
    assert "exportDocx = async function exportDocxLocalBackend" in response.text
    assert "標準テンプレート管理（ローカル）" in response.text
    assert "個人Dify環境には送信されません" in response.text


def test_activate_endpoint_calls_existing_window_callback():
    calls = []
    response = TestClient(create_app(lambda: calls.append("activate"))).post("/api/app/activate")
    assert response.status_code == 204
    assert response.headers[APP_HEADER] == APP_HEADER_VALUE
    assert calls == ["activate"]


def test_existing_instance_requires_health_payload_and_identity_header(monkeypatch):
    class Response:
        headers = {APP_HEADER: APP_HEADER_VALUE}

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return None

        def read(self):
            return b'{"status":"ok"}'

    monkeypatch.setattr(desktop, "_request", lambda *_args, **_kwargs: Response())
    assert desktop.is_existing_instance()
    Response.headers = {}
    assert not desktop.is_existing_instance()


def test_second_instance_activates_without_starting_backend(monkeypatch):
    monkeypatch.setattr(desktop, "acquire_single_instance_mutex", lambda: False)
    monkeypatch.setattr(desktop, "activate_existing_instance", lambda: True)
    monkeypatch.setattr(
        desktop,
        "start_backend",
        lambda: (_ for _ in ()).throw(AssertionError()),
    )
    assert desktop.run_desktop() == 0


def test_template_data_returns_404_before_registration(tmp_path):
    client = TestClient(create_app(data_root=tmp_path))
    response = client.post("/api/template-data", json={"document_type": "server_storage"})
    assert response.status_code == 404
    assert "未登録" in response.json()["detail"]


def test_register_template_and_read_local_template_data(tmp_path):
    client = TestClient(create_app(data_root=tmp_path))
    content = _valid_template_bytes()

    registered = _register_template(client, content)
    assert registered.status_code == 200, registered.text
    body = registered.json()
    assert body["success"] is True
    assert body["document_type"] == "server_storage"
    assert body["template_version"] == "1.2"
    assert body["template_id"]

    loaded = client.post("/api/template-data", json={"document_type": "server_storage"})
    assert loaded.status_code == 200, loaded.text
    data = loaded.json()
    assert data["document_type"] == "server_storage"
    assert data["template_id"] == body["template_id"]
    assert data["template_version"] == "1.2"
    assert isinstance(data["master_json"], list) and data["master_json"]
    assert isinstance(data["chapter_list_json"], list) and data["chapter_list_json"]
    assert isinstance(data["section_contents"], list) and data["section_contents"]
    assert data["returned_section_count"] == len(data["section_contents"])
    assert "概要" in data["reference_text"]

    mirror = tmp_path / "server_storage"
    assert (mirror / "template.docx").read_bytes() == content
    assert (mirror / "metadata.json").exists()
    assert (mirror / "master.json").exists()
    assert (mirror / "chapter_list.json").exists()
    assert (mirror / "section_contents.json").exists()


def test_register_rejects_non_docx(tmp_path):
    client = TestClient(create_app(data_root=tmp_path))
    response = client.post(
        "/api/templates/register",
        data={"document_type": "server_storage"},
        files={"template_file": ("template.txt", b"not-a-docx", "text/plain")},
    )
    assert response.status_code == 400
    assert ".docx" in response.json()["detail"]


def test_generate_word_requires_registered_template(tmp_path):
    client = TestClient(create_app(data_root=tmp_path))
    response = client.post(
        "/api/generate-word",
        json={"document_type": "server_storage", "chapters_json": []},
    )
    assert response.status_code == 400
    assert "未登録" in response.json()["detail"]


def test_generate_word_returns_docx_from_registered_template(tmp_path):
    client = TestClient(create_app(data_root=tmp_path))
    registered = _register_template(client)
    assert registered.status_code == 200, registered.text

    response = client.post(
        "/api/generate-word",
        json={
            "document_type": "server_storage",
            "client_name": "テスト株式会社 様",
            "project_name": "ローカル生成確認",
            "version": "2.0",
            "issue_date": "2026/09/07",
            "project_no": "P-001",
            "revision_history_json": [
                {
                    "issue_date": "2026/09/07",
                    "version": "2.0",
                    "editor": "Tester",
                    "description": "初版",
                }
            ],
            "chapters_json": [
                {
                    "id": "chapter-1",
                    "level": 1,
                    "title": "生成章",
                    "selected": True,
                    "blocks": [
                        {
                            "type": "paragraph",
                            "paragraph_style": "level_0",
                            "text": "ローカルWord生成本文です。",
                        }
                    ],
                    "children": [],
                }
            ],
            "output_filename": "案件_基本設計書.docx",
        },
    )
    assert response.status_code == 200, response.text
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "filename*=UTF-8''" in response.headers["content-disposition"]

    generated = Document(BytesIO(response.content))
    paragraph_texts = [paragraph.text for paragraph in generated.paragraphs]
    assert "テスト株式会社 様" in paragraph_texts
    assert "ローカル生成確認" in paragraph_texts
    assert "生成章" in paragraph_texts
    assert "ローカルWord生成本文です。" in paragraph_texts
