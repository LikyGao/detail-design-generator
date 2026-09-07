from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient

from local_backend.app import create_app
from local_backend.services.staging import remove_staged_file, staged_path


def _template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("株式会社〇〇 様", style="Title")
    document.add_paragraph("サンプルプロジェクト", style="Title")
    document.add_table(rows=4, cols=3)
    document.add_table(rows=2, cols=4)
    document.add_paragraph("1 概要", style="Heading 1")
    document.add_paragraph("標準本文", style="Normal")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def test_staged_word_generation_returns_real_docx_token(tmp_path):
    client = TestClient(create_app(data_root=tmp_path))
    registered = client.post(
        "/api/templates/register",
        data={"document_type": "server_storage", "template_version": "stage-test"},
        files={
            "template_file": (
                "template.docx",
                _template_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert registered.status_code == 200, registered.text

    response = client.post(
        "/api/generate-word/stage",
        json={
            "document_type": "server_storage",
            "client_name": "保存テスト株式会社 様",
            "project_name": "Native Save",
            "chapters_json": [
                {
                    "id": "c1",
                    "level": 1,
                    "title": "保存確認",
                    "selected": True,
                    "blocks": [{"type": "paragraph", "text": "保存本文", "paragraph_style": "level_0"}],
                    "children": [],
                }
            ],
            "output_filename": "Native_Save.docx",
        },
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["filename"] == "Native_Save.docx"
    path = staged_path(body["token"], ".docx")
    try:
        assert path.read_bytes().startswith(b"PK")
        generated = Document(path)
        texts = [paragraph.text for paragraph in generated.paragraphs]
        assert "保存確認" in texts
        assert "保存本文" in texts
    finally:
        remove_staged_file(path)
