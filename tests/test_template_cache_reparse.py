from __future__ import annotations

import json
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from local_backend.services.template_service import (
    LOCAL_TEMPLATE_PARSE_SCHEMA,
    LOCAL_TEMPLATE_PARSE_SCHEMA_FIELD,
    TemplateService,
)
from tools.template_store import _storage_keys


def _style4_template_bytes() -> bytes:
    document = Document()
    document.add_paragraph("株式会社〇〇 様", style="Title")
    document.add_paragraph("サンプルプロジェクト", style="Title")
    document.add_table(rows=4, cols=3)
    document.add_table(rows=2, cols=4)
    document.add_paragraph("1 バックアップ", style="Heading 1")

    style4 = document.styles.add_style("スタイル4", WD_STYLE_TYPE.PARAGRAPH)

    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), "88")
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "4")
    for tag, value in (
        ("w:start", "1"),
        ("w:numFmt", "bullet"),
        ("w:lvlText", "・"),
        ("w:pStyle", style4.style_id),
    ):
        child = OxmlElement(tag)
        child.set(qn("w:val"), value)
        level.append(child)
    p_pr = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "964")
    indent.set(qn("w:hanging"), "170")
    p_pr.append(indent)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), "88")
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), "88")
    num.append(abstract_ref)
    numbering.append(num)

    document.add_paragraph(
        "サーバー障害に備え、システムリストアを目的としたバックアップデータ",
        style=style4,
    )

    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _target_paragraph(data: dict):
    for section in data["section_contents"]:
        for paragraph in section.get("paragraphs") or []:
            if paragraph.get("text", "").startswith("サーバー障害に備え"):
                return paragraph
    raise AssertionError("target paragraph not found")


def test_current_parser_identifies_native_style4(tmp_path):
    service = TemplateService(tmp_path)
    service.register(
        "server_storage",
        "基本設計書_server_v2.1.docx",
        _style4_template_bytes(),
        "2.1",
    )

    paragraph = _target_paragraph(service.get_data("server_storage"))
    # UI スタイル4 is native ilvl 4 and internal paragraph_style level_5.
    assert paragraph["paragraph_style"] == "level_5"
    assert paragraph["native_ilvl"] == 4
    assert paragraph["word_style_name"] == "スタイル4"
    assert paragraph["word_style_id"]


def test_missing_parser_schema_reparses_old_local_cache(tmp_path):
    service = TemplateService(tmp_path)
    service.register(
        "server_storage",
        "基本設計書_server_v2.1.docx",
        _style4_template_bytes(),
        "2.1",
    )

    # Simulate a cache produced by an older local EXE: same DOCX/template id,
    # but stale paragraph metadata and no parser schema marker.
    keys = _storage_keys("server_storage")
    stale = service.get_data("server_storage")["section_contents"]
    paragraph = _target_paragraph({"section_contents": stale})
    paragraph["paragraph_style"] = "level_0"
    paragraph.pop("word_style_id", None)
    paragraph.pop("word_style_name", None)
    paragraph.pop("native_ilvl", None)
    service.storage.set(
        keys["section_contents_key"],
        json.dumps(stale, ensure_ascii=False, separators=(",", ":")).encode("utf-8"),
    )

    metadata_path = tmp_path / "server_storage" / "metadata.json"
    metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    metadata.pop(LOCAL_TEMPLATE_PARSE_SCHEMA_FIELD, None)
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")

    repaired = service.get_data("server_storage")
    assert repaired["template_cache_reparsed"] is True
    assert repaired["local_template_parse_schema"] == LOCAL_TEMPLATE_PARSE_SCHEMA
    repaired_paragraph = _target_paragraph(repaired)
    assert repaired_paragraph["paragraph_style"] == "level_5"
    assert repaired_paragraph["native_ilvl"] == 4
    assert repaired_paragraph["word_style_name"] == "スタイル4"

    refreshed_metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
    assert refreshed_metadata[LOCAL_TEMPLATE_PARSE_SCHEMA_FIELD] == LOCAL_TEMPLATE_PARSE_SCHEMA


def test_current_cache_does_not_reparse_on_every_load(tmp_path):
    service = TemplateService(tmp_path)
    service.register(
        "server_storage",
        "基本設計書_server_v2.1.docx",
        _style4_template_bytes(),
        "2.1",
    )
    first = service.get_data("server_storage")
    second = service.get_data("server_storage")
    assert first["template_cache_reparsed"] is False
    assert second["template_cache_reparsed"] is False
