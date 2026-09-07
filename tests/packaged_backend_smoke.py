"""End-to-end smoke test against the already-running packaged desktop backend."""
from __future__ import annotations

from io import BytesIO

import httpx
from docx import Document

BASE_URL = "http://127.0.0.1:8765"
APP_HEADER = "X-Detail-Design-Generator"
APP_HEADER_VALUE = "local-desktop-v2"


def make_template() -> bytes:
    document = Document()
    document.add_paragraph("株式会社〇〇 様", style="Title")
    document.add_paragraph("サンプルプロジェクト", style="Title")
    document.add_table(rows=4, cols=3)
    document.add_table(rows=2, cols=4)
    document.add_paragraph("1 概要", style="Heading 1")
    document.add_paragraph("概要の標準本文です。", style="Normal")
    document.add_paragraph("1.1 目的", style="Heading 2")
    document.add_paragraph("目的の標準本文です。", style="Normal")
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def require_identity(response: httpx.Response) -> None:
    assert response.headers.get(APP_HEADER) == APP_HEADER_VALUE, response.headers


def word_payload() -> dict:
    return {
        "document_type": "server_storage",
        "client_name": "CIテスト株式会社 様",
        "project_name": "パッケージ版E2E確認",
        "version": "1.0",
        "issue_date": "2026/09/07",
        "project_no": "CI-001",
        "revision_history_json": [
            {
                "issue_date": "2026/09/07",
                "version": "1.0",
                "editor": "CI",
                "description": "E2E smoke",
            }
        ],
        "chapters_json": [
            {
                "id": "chapter-ci",
                "level": 1,
                "title": "E2E生成章",
                "selected": True,
                "blocks": [
                    {
                        "type": "paragraph",
                        "paragraph_style": "level_0",
                        "text": "パッケージ済みEXEから生成した本文です。",
                    }
                ],
                "children": [],
            }
        ],
        "output_filename": "CI_基本設計書.docx",
    }


def main() -> None:
    with httpx.Client(base_url=BASE_URL, timeout=30.0) as client:
        health = client.get("/api/health")
        health.raise_for_status()
        require_identity(health)
        assert health.json() == {"status": "ok"}

        root = client.get("/")
        root.raise_for_status()
        require_identity(root)
        bridge_tag = '<script src="/local-bridge.js"></script>'
        assert bridge_tag in root.text
        bridge_end = root.text.rfind(bridge_tag) + len(bridge_tag)
        body_end = root.text.rfind("</body>")
        assert body_end != -1
        assert bridge_end < body_end
        assert root.text[bridge_end:body_end].strip() == ""

        bridge = client.get("/local-bridge.js")
        bridge.raise_for_status()
        require_identity(bridge)
        for marker in (
            "/api/templates/register",
            "/api/template-data",
            "/api/generate-word/stage",
            "/api/project/stage",
            "/api/preview-state",
            "open_preview_window",
            "ddg-detached-preview",
            "標準テンプレート管理（ローカル）",
        ):
            assert marker in bridge.text, marker

        preview_shell = client.get("/preview-window")
        preview_shell.raise_for_status()
        require_identity(preview_shell)
        assert 'id="standalonePreview"' in preview_shell.text

        preview_write = client.post(
            "/api/preview-state",
            json={"html": "<div>CI preview</div>", "css": "body{}", "revision": 501},
        )
        preview_write.raise_for_status()
        require_identity(preview_write)
        preview_read = client.get("/api/preview-state")
        preview_read.raise_for_status()
        require_identity(preview_read)
        assert preview_read.json()["revision"] == 501
        assert "CI preview" in preview_read.json()["html"]

        project_stage = client.post(
            "/api/project/stage",
            json={
                "project": {
                    "format": "detail-design-generator-project",
                    "schema_version": 1,
                    "phase": "edit",
                    "doc": {"document_type": "server_storage", "chapters": []},
                },
                "suggested_name": "CI.ddgproj",
            },
        )
        project_stage.raise_for_status()
        require_identity(project_stage)
        assert project_stage.json()["token"]
        assert project_stage.json()["filename"] == "CI.ddgproj"

        registered = client.post(
            "/api/templates/register",
            data={"document_type": "server_storage", "template_version": "ci-smoke"},
            files={
                "template_file": (
                    "smoke_template.docx",
                    make_template(),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
        )
        registered.raise_for_status()
        require_identity(registered)
        registration = registered.json()
        assert registration["success"] is True
        assert registration["document_type"] == "server_storage"
        assert registration["template_id"]

        template_data = client.post(
            "/api/template-data", json={"document_type": "server_storage"}
        )
        template_data.raise_for_status()
        require_identity(template_data)
        template_body = template_data.json()
        assert template_body["template_id"] == registration["template_id"]
        assert template_body["master_json"]
        assert template_body["chapter_list_json"]
        assert template_body["section_contents_json"]

        staged_word = client.post("/api/generate-word/stage", json=word_payload())
        staged_word.raise_for_status()
        require_identity(staged_word)
        assert staged_word.json()["token"]
        assert staged_word.json()["filename"] == "CI_基本設計書.docx"

        generated = client.post("/api/generate-word", json=word_payload())
        generated.raise_for_status()
        require_identity(generated)
        assert generated.content.startswith(b"PK")
        assert "filename*=UTF-8''" in generated.headers.get("content-disposition", "")

        output = Document(BytesIO(generated.content))
        paragraphs = [paragraph.text for paragraph in output.paragraphs]
        for expected in (
            "CIテスト株式会社 様",
            "パッケージ版E2E確認",
            "E2E生成章",
            "パッケージ済みEXEから生成した本文です。",
        ):
            assert expected in paragraphs, expected

    print("Packaged desktop backend E2E smoke passed")


if __name__ == "__main__":
    main()
